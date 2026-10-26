import hashlib
import json
import os
import urllib.error
import urllib.request
from io import BytesIO
from pathlib import Path

from database import get_connection
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
import xlrd


def _ollama(path, payload, timeout=None):
    base_url = os.environ.get("OLLAMA_BASE_URL", "http://127.0.0.1:11434").rstrip("/")
    request = urllib.request.Request(
        f"{base_url}{path}", data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout or int(os.environ.get("OLLAMA_TIMEOUT", "300"))) as response:
            return json.loads(response.read())
    except (urllib.error.URLError, TimeoutError) as exc:
        raise RuntimeError(f"No fue posible conectar con Ollama en {base_url}: {exc}") from exc


def _embed(texts):
    result = _ollama("/api/embed", {"model": os.environ.get("OLLAMA_EMBED_MODEL", "nomic-embed-text"), "input": texts})
    return result["embeddings"]


def _chunks(text, size=1200, overlap=200):
    clean = " ".join(text.split())
    step = size - overlap
    return [clean[start:start + size] for start in range(0, len(clean), step) if clean[start:start + size].strip()]


def _vector_literal(values):
    return "[" + ",".join(f"{float(value):.9g}" for value in values) + "]"


ALLOWED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".pptx", ".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml"}


def ingest_pages(document_name, pages, raw=None):
    records = [(page, index, chunk) for page, text in pages for index, chunk in enumerate(_chunks(text))]
    if not records:
        raise ValueError("El manual no contiene texto indexable.")
    digest = hashlib.sha256(raw if raw is not None else "\n".join(text for _, text in pages).encode("utf-8")).hexdigest()
    embeddings = _embed([record[2] for record in records])
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM sop_chunks WHERE document_name=%s", (document_name,))
            for (page, index, content), embedding in zip(records, embeddings):
                cursor.execute(
                    """INSERT INTO sop_chunks
                    (document_hash,document_name,page,chunk_index,content,embedding)
                    VALUES (%s,%s,%s,%s,%s,%s::vector)""",
                    (digest, document_name, page, index, content, _vector_literal(embedding)),
                )
    return len(records)


def ingest_text(document_name, text):
    return ingest_pages(document_name, [(1, text)], text.encode("utf-8"))


def extract_document(filename, raw):
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_DOCUMENT_EXTENSIONS:
        raise ValueError(f"Tipo de archivo no compatible: {extension or 'sin extensión'}")
    stream = BytesIO(raw)
    if extension == ".pdf":
        return [(index + 1, page.extract_text() or "") for index, page in enumerate(PdfReader(stream).pages)]
    if extension == ".docx":
        document = Document(stream)
        return [(1, "\n".join(paragraph.text for paragraph in document.paragraphs))]
    if extension == ".xlsx":
        workbook = load_workbook(stream, read_only=True, data_only=True)
        return [(index + 1, f"Hoja: {sheet.title}\n" + "\n".join(" | ".join(str(value) for value in row if value is not None) for row in sheet.iter_rows(values_only=True))) for index, sheet in enumerate(workbook.worksheets)]
    if extension == ".xls":
        workbook = xlrd.open_workbook(file_contents=raw)
        return [(index + 1, f"Hoja: {sheet.name}\n" + "\n".join(" | ".join(str(sheet.cell_value(row, column)) for column in range(sheet.ncols)) for row in range(sheet.nrows))) for index, sheet in enumerate(workbook.sheets())]
    if extension == ".pptx":
        presentation = Presentation(stream)
        return [(index + 1, "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))) for index, slide in enumerate(presentation.slides)]
    return [(1, raw.decode("utf-8", errors="replace"))]


def ingest_document(filename, raw):
    return ingest_pages(filename, extract_document(filename, raw), raw)


def list_documents():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("""SELECT document_name,count(*) AS chunks,min(created_at) AS indexed_at
                FROM sop_chunks GROUP BY document_name ORDER BY document_name""")
            return cursor.fetchall()


def delete_document(document_name):
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM sop_chunks WHERE document_name=%s RETURNING id", (document_name,))
            return len(cursor.fetchall())


def retrieve(question, limit=3, document_name=None):
    query_vector = _vector_literal(_embed([question])[0])
    where = "WHERE document_name=%s" if document_name else ""
    params = (document_name, query_vector, limit) if document_name else (query_vector, limit)
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                f"""SELECT document_name,page,content,embedding <=> %s::vector AS distance
                FROM sop_chunks {where}
                ORDER BY embedding <=> %s::vector LIMIT %s""" if not document_name else
                f"""SELECT document_name,page,content,embedding <=> %s::vector AS distance
                FROM sop_chunks {where}
                ORDER BY embedding <=> %s::vector LIMIT %s""",
                (query_vector, query_vector, limit) if not document_name else (query_vector, document_name, query_vector, limit),
            )
            return cursor.fetchall()


def answer_rag(question, mode="auto"):
    chunks = retrieve(question)
    if not chunks:
        raise ValueError("No hay manuales indexados. Ejecuta el comando ingest-machine-manuals.")
    evidence = []
    sources = []
    for chunk in chunks:
        label = f"{chunk['document_name']} · página {chunk['page']}"
        evidence.append(f"[{label}]\n{chunk['content']}")
        if label not in sources:
            sources.append(label)
    instruction = (
        "Guía el diagnóstico paso a paso. Prioriza parada segura, aislamiento de energías, LOTO, inocuidad y escalamiento. "
        if mode == "troubleshooting" else "Responde de forma precisa usando solo la evidencia recuperada. "
    )
    prompt = f"""{instruction}
No inventes datos. No autorices retirar guardas, anular protecciones ni intervenir equipos energizados. Responde en texto plano, sin Markdown.
Si la evidencia es insuficiente, indícalo. Cita las fuentes entre corchetes.
Responde en español en máximo 180 palabras.

EVIDENCIA:
{chr(10).join(evidence)}

PREGUNTA: {question}"""
    result = _ollama("/api/chat", {
        "model": os.environ.get("OLLAMA_CHAT_MODEL", "llama3.1:8b"),
        "messages": [
            {"role": "system", "content": "Eres un copiloto documental para mantenimiento y producción de una planta de alimentos. No sustituyes decisiones de personal autorizado."},
            {"role": "user", "content": prompt},
        ],
        "stream": False, "keep_alive": "10m",
        "options": {"temperature": 0.1, "num_ctx": 4096, "num_predict": int(os.environ.get("OLLAMA_MAX_TOKENS", "220"))},
    })
    return result["message"]["content"], sources


def resolve_mode(question, requested="auto"):
    if requested != "auto":
        return requested
    normalized = question.casefold()
    predictive_terms = ("riesgo", "predictiv", "máquina", "maquina", "vibración", "vibracion", "temperatura", "corriente", "rpm", "presión", "presion", "critical", "warning")
    troubleshooting_terms = ("alarma", "falla", "falló", "fallo", "detenida", "no arranca", "atasco", "diagnóstico", "diagnostico")
    if any(term in normalized for term in predictive_terms):
        return "predictive"
    if any(term in normalized for term in troubleshooting_terms):
        return "troubleshooting"
    return "manuals"


def answer_predictive(question):
    from predictive import predictive_payload

    payload = predictive_payload(refresh=False)
    if not payload["machines"]:
        raise ValueError("No existen lecturas predictivas disponibles.")
    normalized = question.casefold()
    selected = [machine for machine in payload["machines"] if machine["id"].casefold() in normalized or machine["name"].casefold() in normalized]
    rows = selected[:4] or payload["machines"][:5]
    context = [
        f"Resumen: {payload['summary']['healthy']} saludables, {payload['summary']['warning']} en alerta, {payload['summary']['critical']} críticas; riesgo promedio {payload['summary']['average_risk']}%.",
        f"Última captura: {payload['generated_at']}. El índice es demostrativo y no una probabilidad estadística validada.",
    ]
    for machine in rows:
        factors = ", ".join(f"{factor['label']} {factor['contribution']} puntos" for factor in machine["factors"][:2])
        limits = machine["limits"]
        context.append(f"{machine['id']} {machine['name']}: riesgo {machine['risk']}%, estado {machine['status']}, vibración {machine['vibration']} mm/s (límite {limits['vibration']}), temperatura {machine['temperature']} °C (límite {limits['temperature']}), corriente {machine['current']} A (nominal {limits['current']}), RPM {machine['rpm']} (nominal {limits['nominal_rpm']}), presión {machine['pressure']} bar (máxima {limits['pressure']}), uso {machine['usage_hours']} h (máximo {limits['usage_hours']}). Factores calculados: {factors}.")
    manual_chunks = retrieve(question, limit=3)
    sources = [f"PostgreSQL · machine_readings · {payload['generated_at']}"]
    evidence = []
    for chunk in manual_chunks:
        label = f"{chunk['document_name']} · página {chunk['page']}"
        evidence.append(f"[{label}] {chunk['content']}")
        if label not in sources:
            sources.append(label)
    prompt = f"""Responde únicamente con los datos y límites suministrados. Distingue lectura medida, comparación con límite, interpretación y acción recomendada.
Está prohibido inventar, recalcular o mencionar cualquier número que no aparezca en DATOS ACTUALES. No denomines "límite" al riesgo promedio ni a la contribución de un factor.
Si un valor o límite no está disponible, di que no está disponible. No autorices intervenir sin parada segura/LOTO. Recomienda validación humana.
Usa texto plano sin Markdown, máximo 220 palabras, oraciones breves y termina completamente la última oración.

DATOS ACTUALES:
{chr(10).join(context)}

MANUALES RECUPERADOS:
{chr(10).join(evidence)}

PREGUNTA: {question}"""
    result = _ollama("/api/chat", {
        "model": os.environ.get("OLLAMA_CHAT_MODEL", "llama3.1:8b"),
        "messages": [{"role": "system", "content": "Eres un copiloto de mantenimiento predictivo para una planta de alimentos. Debes ser fiel a los datos proporcionados y nunca fabricar cifras o límites."}, {"role": "user", "content": prompt}],
        "stream": False, "keep_alive": "10m",
        "options": {"temperature": 0.0, "num_ctx": 4096, "num_predict": int(os.environ.get("PREDICTIVE_CHAT_MAX_TOKENS", "360"))},
    })
    return result["message"]["content"], sources


def answer_chat(question, requested_mode="auto"):
    mode = resolve_mode(question, requested_mode)
    if mode == "predictive":
        answer, sources = answer_predictive(question)
    else:
        answer, sources = answer_rag(question, mode)
    return answer, sources, mode
